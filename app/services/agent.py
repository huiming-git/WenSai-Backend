import json
import logging
import os
import re

from openai import OpenAI
from PyPDF2 import PdfReader
from pptx import Presentation

from app.config import LLM_API_KEY, LLM_BASE_URL, LLM_MODEL

SYSTEM_PROMPT = """You are a senior academic paper reviewer.
Review the given paper and respond STRICTLY in the following JSON format (no other text):

{
  "score": <integer 1-10>,
  "recommendation": "<one of: accept, minor_revision, major_revision, reject>",
  "content": "<your detailed review in markdown format, including: Summary, Strengths, Weaknesses, Detailed Comments, and Recommendation>"
}

Scoring guide:
- 9-10: Excellent, publishable as-is
- 7-8: Good, minor revisions needed
- 5-6: Average, major revisions needed
- 3-4: Below average, fundamental issues
- 1-2: Poor, recommend rejection

Write the review content in Chinese (中文)."""

MAX_TEXT_CHARS = 80000  # ~20K tokens, safe limit for most models

logger = logging.getLogger(__name__)


def _extract_text(file_path: str) -> str | None:
    """Extract text content from a file."""
    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages).strip()

        elif ext in (".pptx", ".ppt"):
            prs = Presentation(file_path)
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t:
                                texts.append(t)
                if texts:
                    slides.append(f"[Slide {i}]\n" + "\n".join(texts))
            return "\n\n".join(slides).strip()

        elif ext in (".txt", ".md", ".tex"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().strip()

    except Exception:
        pass

    return None


def _build_user_message(title: str, requirements: str | None = None) -> str:
    """Build user message with paper metadata."""
    parts = [f"请评审以下论文：\n\n标题：{title}"]
    if requirements:
        parts.append(f"评审要求：{requirements}")
    return "\n\n".join(parts)


def _append_extracted_text(user_message: str, file_path: str) -> str:
    """Append extracted text from local file if available."""
    if not os.path.exists(file_path):
        return user_message

    text = _extract_text(file_path)
    if not text:
        return user_message

    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n\n[... 正文过长，已截断 ...]"

    return f"{user_message}\n\n论文全文：\n{text}"


def _review_with_chat_completions(client: OpenAI, user_message: str) -> dict:
    """Fallback path: use chat.completions with plain text content."""
    response = client.chat.completions.create(
        model=LLM_MODEL,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_message},
        ],
        temperature=0.3,
    )
    return _parse_response(response.choices[0].message.content.strip())


def _review_with_responses_file(client: OpenAI, user_message: str, file_path: str) -> dict:
    """Preferred path: upload file and pass file_id to responses API."""
    uploaded_file_id = None
    filename = os.path.basename(file_path)

    try:
        with open(file_path, "rb") as f:
            uploaded = client.files.create(file=f, purpose="user_data")
        uploaded_file_id = uploaded.id

        response = client.responses.create(
            model=LLM_MODEL,
            instructions=SYSTEM_PROMPT,
            input=[
                {
                    "role": "user",
                    "content": [
                        {"type": "input_text", "text": user_message},
                        {"type": "input_file", "file_id": uploaded_file_id, "filename": filename},
                    ],
                }
            ],
            temperature=0.3,
        )
        result_text = (response.output_text or "").strip()
        if not result_text:
            raise ValueError("LLM response text is empty")
        return _parse_response(result_text)
    finally:
        if uploaded_file_id:
            try:
                client.files.delete(uploaded_file_id)
            except Exception:
                # Non-fatal cleanup failure.
                logger.warning("Failed to delete uploaded file from LLM provider: %s", uploaded_file_id)


def _parse_response(result_text: str) -> dict:
    """Parse LLM response into {score, content, recommendation}."""
    try:
        json_match = re.search(r'```(?:json)?\s*([\s\S]*?)\s*```', result_text)
        if json_match:
            result_text = json_match.group(1)

        data = json.loads(result_text)
        return {
            "score": max(1, min(10, int(data.get("score", 5)))),
            "content": data.get("content", result_text),
            "recommendation": data.get("recommendation", "minor_revision"),
        }
    except (json.JSONDecodeError, KeyError):
        return {
            "score": 5,
            "content": result_text,
            "recommendation": "minor_revision",
        }


def _build_llm_log(**kwargs) -> str:
    """Serialize LLM execution diagnostics as JSON string."""
    return json.dumps(kwargs, ensure_ascii=False)


def review_paper(title: str, requirements: str | None = None, file_path: str | None = None) -> dict:
    """Call LLM to review a paper.

    Preferred path: upload file and send file_id via Responses API.
    Fallback path: extract text locally and use chat.completions.
    """
    if not LLM_API_KEY or LLM_API_KEY == "sk-xxx":
        raise ValueError("LLM_API_KEY not configured. Set it in backend/.env")

    client = OpenAI(api_key=LLM_API_KEY, base_url=LLM_BASE_URL)
    user_message = _build_user_message(title, requirements)

    if file_path and os.path.exists(file_path):
        try:
            result = _review_with_responses_file(client, user_message, file_path)
            result["llm_log"] = _build_llm_log(
                mode="responses_input_file",
                model=LLM_MODEL,
                base_url=LLM_BASE_URL,
                file_name=os.path.basename(file_path),
            )
            return result
        except Exception as e:
            logger.warning(
                "Responses file path failed (%s), fallback to extracted text + chat.completions.",
                str(e),
            )
            message_with_text = _append_extracted_text(user_message, file_path)
            text_extracted = message_with_text != user_message
            result = _review_with_chat_completions(client, message_with_text)
            result["llm_log"] = _build_llm_log(
                mode="chat_completions_fallback",
                fallback_from="responses_input_file",
                fallback_reason=str(e)[:500],
                text_extracted=text_extracted,
                model=LLM_MODEL,
                base_url=LLM_BASE_URL,
                file_name=os.path.basename(file_path),
            )
            return result

    result = _review_with_chat_completions(client, user_message)
    result["llm_log"] = _build_llm_log(
        mode="chat_completions_no_file",
        model=LLM_MODEL,
        base_url=LLM_BASE_URL,
        file_name=os.path.basename(file_path) if file_path else None,
        file_exists=bool(file_path and os.path.exists(file_path)),
    )
    return result
